import base64
import datetime
import io
import os
import re
import shutil
import zipfile

import openpyxl
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

# --- SUPPORT ROBUSTE DES PHOTOS HEIC (Samsung S24/S25/S26, etc.) ---
try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:
    pass

st.set_page_config(
    page_title="Frais_fabrix - Gestion des Notes de Frais",
    page_icon="💶",
    layout="wide",
)

# --- CHARTE GRAPHIQUE PROFESSIONNELLE & ADAPTATION MOBILE 9:16 ---
st.markdown(
    """
    <style>
        /* Optimisation globale pour container principal et affichage mobile */
        .block-container {
            padding-top: 1.5rem !important;
            padding-bottom: 2rem !important;
            padding-left: 1rem !important;
            padding-right: 1rem !important;
            max-width: 100% !important;
        }

        /* Fond global de la barre latérale (Vert très foncé Cleemy) */
        [data-testid="stSidebar"] {
            background-color: #18302c !important;
            padding-top: 1rem;
        }

        /* Forcer tous les textes, labels, paragraphes et titres de la sidebar en blanc éclatant */
        [data-testid="stSidebar"] *, 
        [data-testid="stSidebar"] label, 
        [data-testid="stSidebar"] span, 
        [data-testid="stSidebar"] p, 
        [data-testid="stSidebar"] div,
        [data-testid="stSidebar"] .stMarkdown {
            color: #ffffff !important;
        }

        /* Style des boutons dans la sidebar (Vert Cleemy professionnel) */
        [data-testid="stSidebar"] .stButton button {
            background-color: #229954 !important;
            color: #ffffff !important;
            border: none !important;
            border-radius: 6px !important;
            font-weight: 600 !important;
            width: 100%;
        }
        
        [data-testid="stSidebar"] .stButton button:hover {
            background-color: #1e8449 !important;
            color: #ffffff !important;
        }

        /* Style des menus déroulants (selectbox) dans la sidebar */
        [data-testid="stSidebar"] .stSelectbox div[data-baseweb="select"] {
            background-color: #122421 !important;
            color: #ffffff !important;
            border-color: #2b4d47 !important;
        }
        
        [data-testid="stSidebar"] .stSelectbox svg {
            fill: #ffffff !important;
        }

        /* --- AMÉLIORATION DES CHAMPS DE FORMULAIRE (Texte noir & relief net) --- */
        div[data-testid="stForm"] input, 
        div[data-testid="stForm"] div[data-baseweb="select"] div,
        div[data-testid="stForm"] .stDateInput input {
            background-color: #ffffff !important;
            color: #000000 !important;
            border: 1px solid #cbd5e1 !important;
            border-radius: 6px !important;
            box-shadow: 0 1px 2px rgba(0, 0, 0, 0.05) !important;
        }

        div[data-testid="stForm"] input, 
        div[data-testid="stForm"] .stDateInput input {
            color: #000000 !important;
        }

        div[data-testid="stForm"] input:focus, 
        div[data-testid="stForm"] div[data-baseweb="select"]:focus-within,
        div[data-testid="stForm"] .stDateInput input:focus {
            border-color: #229954 !important;
            box-shadow: 0 0 0 3px rgba(34, 153, 84, 0.15) !important;
        }

        /* --- PERSONNALISATION DES BOUTONS PRIMAIRES (STYLE CLEEMY VERT) --- */
        button[kind="primary"] {
            background-color: #229954 !important;
            border-color: #229954 !important;
            color: #ffffff !important;
        }
        button[kind="primary"]:hover {
            background-color: #1e8449 !important;
            border-color: #1e8449 !important;
        }

        /* --- RÉDUCTION DE LA LARGEUR DES CHAMPS ET DE L'UPLOADER SUR PC --- */
        @media (min-width: 769px) {
            div[data-testid="stForm"] input, 
            div[data-testid="stForm"] div[data-baseweb="select"],
            div[data-testid="stForm"] .stDateInput input,
            div[data-testid="stFileUploader"] {
                max-width: 380px !important;
            }
        }

        .stButton button {
            border-radius: 6px !important;
            font-weight: 600 !important;
            transition: all 0.2s ease-in-out;
        }

        /* --- OPTIMISATION ÉCRAN SMARTPHONE & FORMAT VERTICAL (9:16) --- */
        @media (max-width: 768px) {
            /* Forcer les colonnes principales à s'empiler proprement sur mobile */
            [data-testid="stHorizontalBlock"] {
                flex-direction: column !important;
            }
            [data-testid="column"] {
                width: 100% !important;
                flex: 1 1 100% !important;
                min-width: unset !important;
            }
            h1 { font-size: 1.25rem !important; }
            .stSubheader { font-size: 1rem !important; }
            .stButton button {
                width: 100% !important;
            }
            /* Réduction de la taille des chiffres dans st.metric sur smartphone */
            [data-testid="stMetricValue"] {
                font-size: 1.15rem !important;
            }
            [data-testid="stMetricLabel"] {
                font-size: 0.75rem !important;
            }
        }
    </style>
    """,
    unsafe_allow_html=True,
)


def _mdp(nom: str) -> str:
    try:
        return st.secrets["passwords"][nom]
    except Exception:
        return "123"


UTILISATEURS = {
    "Sabrina": {"nom_complet": "Sabrina", "password": _mdp("Sabrina")},
    "Alain": {"nom_complet": "Alain Autier", "password": _mdp("Alain")},
    "Edouard": {"nom_complet": "Edouard Dupont", "password": _mdp("Alain")},
}

COLS_MONTANTS = [
    "Parking",
    "Entretien matériel",
    "Réception",
    "Hôtel",
    "Restauration",
    "Gasoil",
    "TVA",
    "Total",
]


# --- VISUALISEUR INTERACTIF AVEC PINCH-TO-ZOOM TACTILE ---
def afficher_image_zoomable(img_b64: str, hauteur: int = 420):
    html = """
    <div id="v" style="width:100%; height:__H__px; overflow:hidden; position:relative;
            touch-action:none; border:1px solid #ddd; border-radius:8px; background:#111; display:flex; align-items:center; justify-content:center;">
      <img id="im" draggable="false"
            src="data:image/jpeg;base64,__IMG__"
            style="max-width:100%; max-height:100%; object-fit:contain; cursor:grab; user-select:none; -webkit-user-drag:none; transform-origin:center center; transition: transform 0.05s ease-out;">
    </div>
    <p style="font-size: 0.7rem; color: #666; margin-top: 4px; text-align: center;">💡 <i>Pincez pour zoomer, glissez pour déplacer, double-tap pour réinitialiser.</i></p>
    <script>
      const v = document.getElementById('v'), im = document.getElementById('im');
      let s = 1, x = 0, y = 0;
      const pts = new Map();
      let lastTap = 0;

      function apply() { 
        im.style.transform = `translate(${x}px, ${y}px) scale(${s})`; 
      }

      function zoomAt(cx, cy, factor) {
        const s2 = Math.min(8, Math.max(1, s * factor));
        x = cx - (cx - x) * (s2 / s);
        y = cy - (cy - y) * (s2 / s);
        s = s2;
        if (s === 1) { x = 0; y = 0; }
      }

      function info() {
        const a = [...pts.values()];
        const r = v.getBoundingClientRect();
        return {
          d: Math.hypot(a[0].x - a[1].x, a[0].y - a[1].y),
          mx: (a[0].x + a[1].x) / 2 - r.left,
          my: (a[0].y + a[1].y) / 2 - r.top
        };
      }

      v.addEventListener('pointerdown', e => {
        v.setPointerCapture(e.pointerId);
        pts.set(e.pointerId, {x: e.clientX, y: e.clientY});
        const now = Date.now();
        if (pts.size === 1 && now - lastTap < 300) { s = 1; x = 0; y = 0; apply(); }
        if (pts.size === 1) lastTap = now;
      });

      v.addEventListener('pointermove', e => {
        if (!pts.has(e.pointerId)) return;
        const prev = pts.get(e.pointerId);
        if (pts.size === 2) {
          const before = info();
          pts.set(e.pointerId, {x: e.clientX, y: e.clientY});
          const after = info();
          zoomAt(after.mx, after.my, after.d / before.d);
          x += after.mx - before.mx;
          y += after.my - before.my;
        } else {
          pts.set(e.pointerId, {x: e.clientX, y: e.clientY});
          if (s > 1) { x += e.clientX - prev.x; y += e.clientY - prev.y; }
        }
        apply();
      });

      ['pointerup', 'pointercancel', 'pointerleave'].forEach(ev =>
        v.addEventListener(ev, e => pts.delete(e.pointerId)));

      v.addEventListener('wheel', e => {
        e.preventDefault();
        const r = v.getBoundingClientRect();
        zoomAt(e.clientX - r.left, e.clientY - r.top, e.deltaY < 0 ? 1.15 : 0.87);
        apply();
      }, {passive: false});
    </script>
    """
    html = (
        html.replace("__IMG__", img_b64)
        .replace("__H__", str(hauteur))
        .replace("$", "$")
    )
    components.html(html, height=hauteur + 40)


# --- AUTHENTIFICATION ---
if "authentifie" not in st.session_state:
    st.session_state["authentifie"] = False
    st.session_state["user_key"] = ""
    st.session_state["user_nom_complet"] = ""

if not st.session_state["authentifie"]:
    st.title("🔐 Connexion à Frais_fabrix")
    st.markdown("Veuillez vous identifier pour accéder à vos notes de frais.")

    with st.form("form_login"):
        username_input = st.text_input("Identifiant (Sabrina ou Alain)")
        password_input = st.text_input("Mot de passe", type="password")
        submit_login = st.form_submit_button("Se connecter")

        if submit_login:
            user_trouve = None
            for u in UTILISATEURS:
                if u.lower() == username_input.strip().lower():
                    user_trouve = u
                    break

            if user_trouve and UTILISATEURS[user_trouve]["password"] == password_input:
                st.session_state["authentifie"] = True
                st.session_state["user_key"] = user_trouve.lower()
                st.session_state["user_nom_complet"] = UTILISATEURS[user_trouve][
                    "nom_complet"
                ]
                st.rerun()
            else:
                st.error("Identifiant ou mot de passe incorrect.")
    st.stop()

# --- APPLICATION PRINCIPALE ---
user_key = st.session_state["user_key"]
nom_salarie_defaut = st.session_state["user_nom_complet"]

FICHIER_SAUVEGARDE = f"frais_sauvegarde_{user_key}.csv"
DOSSIER_ARCHIVES = f"archives_frais_{user_key}"
DOSSIER_JUSTIFICATIFS = f"justificatifs_sauvegardés_{user_key}"
os.makedirs(DOSSIER_ARCHIVES, exist_ok=True)
os.makedirs(DOSSIER_JUSTIFICATIFS, exist_ok=True)

colonnes_attendues = [
    "ID",
    "Période",
    "Nom",
    "Date",
    "Libellé",
    "Chantier",
    "Parking",
    "Entretien matériel",
    "Réception",
    "Hôtel",
    "Restauration",
    "Gasoil",
    "TVA",
    "Total",
    "Justificatif",
]

if "frais_data" not in st.session_state:
    if os.path.exists(FICHIER_SAUVEGARDE):
        try:
            st.session_state["frais_data"] = pd.read_csv(FICHIER_SAUVEGARDE)
            df_init = st.session_state["frais_data"]
            if not df_init.empty and "ID" in df_init.columns:
                st.session_state["id_counter"] = int(df_init["ID"].max())
            else:
                st.session_state["id_counter"] = len(df_init)
        except Exception:
            st.session_state["frais_data"] = pd.DataFrame(columns=colonnes_attendues)
            st.session_state["id_counter"] = 0
    else:
        st.session_state["frais_data"] = pd.DataFrame(columns=colonnes_attendues)
        st.session_state["id_counter"] = 0


def sauvegarder_disque():
    st.session_state["frais_data"].to_csv(FICHIER_SAUVEGARDE, index=False)


# --- BARRE LATÉRALE : Gestion enrichie des Archives ---
with st.sidebar:
    st.write(f"👤 Connecté en tant que : **{nom_salarie_defaut}**")
    if st.button("🚪 Se déconnecter"):
        for k in ("authentifie", "user_key", "user_nom_complet"):
            st.session_state[k] = False if k == "authentifie" else ""
        st.session_state.pop("frais_data", None)
        st.session_state.pop("ticket_actuel_bytes", None)
        st.session_state.pop("ticket_actuel_nom", None)
        st.rerun()
    st.markdown("---")

    st.header("📂 Gestion des Mois")
    fichiers_archives = [
        f for f in os.listdir(DOSSIER_ARCHIVES) if f.endswith(".csv")
    ]

    if fichiers_archives:
        st.subheader("Recharger un ancien mois")
        options_archives_dict = {"-- Mois en cours --": None}
        for f in sorted(fichiers_archives, reverse=True):
            chemin_f = os.path.join(DOSSIER_ARCHIVES, f)
            try:
                df_arc = pd.read_csv(chemin_f)
                nb_dep = len(df_arc)
                tot_arc = (
                    df_arc["Total"].sum()
                    if "Total" in df_arc.columns and not df_arc.empty
                    else 0.0
                )
                libelle_enrichi = f"{f} ➔ ({nb_dep} dépense{'s' if nb_dep > 1 else ''} | {tot_arc:.2f} €)"
                options_archives_dict[libelle_enrichi] = f
            except Exception:
                options_archives_dict[f] = f

        archive_choisie_libelle = st.selectbox(
            "Sélectionnez une archive", list(options_archives_dict.keys())
        )
        if archive_choisie_libelle != "-- Mois en cours --":
            if st.sidebar.button("📂 Charger cette archive"):
                nom_fic_archive = options_archives_dict[archive_choisie_libelle]
                chemin_archive = os.path.join(DOSSIER_ARCHIVES, nom_fic_archive)
                st.session_state["frais_data"] = pd.read_csv(chemin_archive)
                st.rerun()

# --- EN-TÊTE PROFESSIONNEL ---
st.markdown("---")
st.title("📄 Frais_fabrix - Suivi des Frais")
st.markdown("---")

# --- INITIALISATION DE LA CLÉ DE L'UPLOADER ---
if "uploader_key" not in st.session_state:
    st.session_state["uploader_key"] = 0

# --- ÉTAPE 1 : UPLOAD DU JUSTIFICATIF & BOUTON NOUVEAU TICKET ---
st.subheader("📸 1. Joignez ou scannez votre ticket")

col_upl1, col_upl2 = st.columns([1.1, 1.9])
with col_upl1:
    photo_uploadee = st.file_uploader(
        "Sélectionnez ou prenez la photo",
        type=["png", "jpg", "jpeg", "heic", "HEIC"],
        key=f"uploader_ticket_{st.session_state['uploader_key']}",
    )
with col_upl2:
    st.write("")
    if st.button("🔄 Nouveau ticket", type="secondary"):
        st.session_state.pop("ticket_actuel_bytes", None)
        st.session_state.pop("ticket_actuel_nom", None)
        st.session_state["uploader_key"] += 1
        st.rerun()

# --- SÉCURISATION DE LA CAPTURE DU FICHIER ---
if photo_uploadee is not None:
    try:
        bytes_data = photo_uploadee.getvalue()
        if bytes_data:
            img_test = Image.open(io.BytesIO(bytes_data))
            img_test.verify()
            st.session_state["ticket_actuel_bytes"] = bytes_data
            st.session_state["ticket_actuel_nom"] = getattr(
                photo_uploadee, "name", "ticket.jpg"
            )
    except Exception:
        pass

nom_fichier_sauvegarde = "Aucun justificatif"
image_rgb = None
angle_rotation = 0

# --- DISPOSITION EN COLONNES (EMPILÉES AUTOMATIQUEMENT SUR MOBILE) ---
col_form, col_ticket = st.columns([1.2, 1], gap="large")

with col_ticket:
    st.subheader("👁️ Visualisation")

    angle_rotation = st.selectbox(
        "🔄 Pivoter l'image", [0, 90, 180, 270], index=0, key="select_rotation"
    )

    if "ticket_actuel_bytes" in st.session_state and st.session_state[
        "ticket_actuel_bytes"
    ]:
        try:
            image_pil = Image.open(io.BytesIO(st.session_state["ticket_actuel_bytes"]))
            image_pil.thumbnail((1600, 1600))

            if angle_rotation > 0:
                image_pil = image_pil.rotate(-angle_rotation, expand=True)

            image_rgb = image_pil.convert("RGB")

            nom_original = st.session_state["ticket_actuel_nom"]
            timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            nom_fichier_sauvegarde = f"ticket_{timestamp_str}.jpg"
            chemin_image_disque = os.path.join(
                DOSSIER_JUSTIFICATIFS, nom_fichier_sauvegarde
            )
            image_rgb.save(chemin_image_disque, format="JPEG", quality=90)
        except Exception as e:
            st.error(f"Erreur de lecture de l'image : {e}")

    if image_rgb is not None:
        buffered = io.BytesIO()
        image_rgb.save(buffered, format="JPEG", quality=90)
        img_str = base64.b64encode(buffered.getvalue()).decode()
        afficher_image_zoomable(img_str, hauteur=400)

        buf_img = io.BytesIO()
        image_rgb.save(buf_img, format="JPEG", quality=90)
        buf_img.seek(0)
        st.download_button(
            label="📥 Télécharger ce justificatif HD",
            data=buf_img,
            file_name=nom_fichier_sauvegarde,
            mime="application/jpeg",
            type="secondary",
        )
    else:
        st.info("👈 Chargez un ticket ci-dessus pour l'afficher ici.")

with col_form:
    st.subheader("✍️ 2. Saisie de la dépense")

    with st.form("form_frais", clear_on_submit=True):
        col_f1, col_f2 = st.columns(2)

        with col_f1:
            periode = st.text_input("Période (ex: août-26)", value="août-26")
            nom = st.text_input("Nom du salarié", value=nom_salarie_defaut)
            date_frais = st.date_input("Date de la dépense", value=datetime.date.today())
            libelle = st.text_input("Libellé / Type", value="")

        with col_f2:
            categorie = st.selectbox(
                "Catégorie",
                [
                    "Parking / Péage",
                    "Restauration",
                    "Gasoil",
                    "Hôtel",
                    "Entretien matériel",
                    "Réception",
                ],
            )
            chantier = st.text_input("Chantier", value="")
            montant_str = st.text_input("Montant TTC (€)", value="0.00")
            tva_str = st.text_input("TVA (€)", value="0.00")

        submitted = st.form_submit_button("💾 Enregistrer la dépense")

        if submitted:
            try:
                montant = float(str(montant_str).replace(",", "."))
                tva = float(str(tva_str).replace(",", "."))
            except ValueError:
                montant = 0.0
                tva = 0.0

            if montant > 0:
                st.session_state["id_counter"] += 1

                nouvelle_ligne = {
                    "ID": st.session_state["id_counter"],
                    "Période": periode,
                    "Nom": nom,
                    "Date": date_frais.strftime("%d-%b"),
                    "Libellé": libelle if libelle else categorie,
                    "Chantier": chantier,
                    "Parking": montant if categorie == "Parking / Péage" else 0.0,
                    "Entretien matériel": (
                        montant if categorie == "Entretien matériel" else 0.0
                    ),
                    "Réception": montant if categorie == "Réception" else 0.0,
                    "Hôtel": montant if categorie == "Hôtel" else 0.0,
                    "Restauration": montant if categorie == "Restauration" else 0.0,
                    "Gasoil": montant if categorie == "Gasoil" else 0.0,
                    "TVA": max(tva, 0.0),
                    "Total": montant,
                    "Justificatif": nom_fichier_sauvegarde,
                }

                st.session_state["frais_data"] = pd.concat(
                    [st.session_state["frais_data"], pd.DataFrame([nouvelle_ligne])],
                    ignore_index=True,
                )
                sauvegarder_disque()

                st.session_state.pop("ticket_actuel_bytes", None)
                st.session_state.pop("ticket_actuel_nom", None)
                st.session_state["uploader_key"] += 1

                st.success("Dépense enregistrée !")
                st.rerun()
            else:
                st.error("Veuillez indiquer un montant supérieur à 0.")

# --- TABLEAU RÉCAPITULATIF & KPIS ---
st.markdown("---")
nb_lignes = len(st.session_state["frais_data"])
texte_compteur = (
    f" ({nb_lignes} dépense{'s' if nb_lignes > 1 else ''})"
    if nb_lignes > 0
    else " (Aucune dépense)"
)
st.subheader(f"📊 Tableau récapitulatif{texte_compteur}")

if not st.session_state["frais_data"].empty:
    df_actuel = st.session_state["frais_data"]
    tot_parking = df_actuel["Parking"].sum()
    tot_restauration = df_actuel["Restauration"].sum()
    tot_gasoil = df_actuel["Gasoil"].sum()
    tot_hotel = df_actuel["Hôtel"].sum()
    tot_entretien = df_actuel["Entretien matériel"].sum()
    tot_reception = df_actuel["Réception"].sum()
    total_general = df_actuel["Total"].sum()

    kpi_col1, kpi_col2, kpi_col3, kpi_col4 = st.columns(4)
    with kpi_col1:
        st.metric(label="💶 Total Général", value=f"{total_general:.2f} €")
    with kpi_col2:
        st.metric(label="🍽 Restauration", value=f"{tot_restauration:.2f} €")
    with kpi_col3:
        st.metric(label="⛽ Gasoil", value=f"{tot_gasoil:.2f} €")
    with kpi_col4:
        st.metric(label="🅿️ Parking", value=f"{tot_parking:.2f} €")

    with st.expander("Voir le détail des autres catégories"):
        kpi_sub1, kpi_sub2, kpi_sub3 = st.columns(3)
        with kpi_sub1:
            st.metric(label="🏨 Hôtel", value=f"{tot_hotel:.2f} €")
        with kpi_sub2:
            st.metric(label="🔧 Entretien", value=f"{tot_entretien:.2f} €")
        with kpi_sub3:
            st.metric(label="🥂 Réception", value=f"{tot_reception:.2f} €")

    st.markdown("---")

    df_travail = st.session_state["frais_data"].copy()
    if "❌ Suppr" not in df_travail.columns:
        df_travail.insert(0, "❌ Suppr", False)

    df_affiche = df_travail.drop(columns=["ID"])
    df_modifie = st.data_editor(
        df_affiche, use_container_width=True, key="tableau_edition"
    )

    colonnes_maj = [c for c in df_modifie.columns if c != "❌ Suppr"]
    st.session_state["frais_data"][colonnes_maj] = df_modifie[colonnes_maj]
    sauvegarder_disque()

    st.markdown("---")
    col_bas1, col_bas2 = st.columns(2)

    with col_bas1:
        st.subheader("🗑️ Actions")
        if st.button("🗑️ Supprimer les lignes cochées"):
            lignes_a_cocher = df_modifie["❌ Suppr"] == True
            if lignes_a_cocher.any():
                fichiers_a_supprimer = st.session_state["frais_data"].loc[
                    lignes_a_cocher, "Justificatif"
                ]
                for nom_fic in fichiers_a_supprimer:
                    if pd.notna(nom_fic) and nom_fic != "Aucun justificatif":
                        chemin_img = os.path.join(DOSSIER_JUSTIFICATIFS, nom_fic)
                        if os.path.exists(chemin_img):
                            try:
                                os.remove(chemin_img)
                            except Exception:
                                pass

                st.session_state["frais_data"] = (
                    st.session_state["frais_data"]
                    .loc[~lignes_a_cocher]
                    .reset_index(drop=True)
                )
                sauvegarder_disque()
                st.success("Supprimé avec succès !")
                st.rerun()
            else:
                st.warning("Cochez au moins une case ❌ Suppr.")

    with col_bas2:
        st.subheader("📁 Déclaration")
        if st.button("📄 Préparer les fichiers Excel/ZIP", type="primary"):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Note de Frais"

            font_titre = Font(name="Arial", size=14, bold=True, color="1F4E78")
            font_sous_titre = Font(name="Arial", size=11, bold=True, color="333333")
            font_en_tete = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            fill_en_tete = PatternFill(
                start_color="1F4E78", end_color="1F4E78", fill_type="solid"
            )
            font_donnees = Font(name="Arial", size=10)
            font_total = Font(name="Arial", size=10, bold=True)
            align_centre = Alignment(horizontal="center", vertical="center")
            align_droite = Alignment(horizontal="right", vertical="center")
            align_gauche = Alignment(horizontal="left", vertical="center")
            gris = Side(style="thin", color="D9D9D9")
            bordure_fine = Border(left=gris, right=gris, top=gris, bottom=gris)
            bordure_total = Border(
                top=Side(style="thin", color="000000"),
                bottom=Side(style="double", color="000000"),
            )

            periode_rapport = (
                str(df_modifie["Période"].iloc[0])
                if "Période" in df_modifie.columns and not df_modifie.empty
                else "Mois"
            )

            ws["A1"] = f"NOTE DE FRAIS - {nom_salarie_defaut.upper()}"
            ws["A1"].font = font_titre
            ws["A2"] = f"Salarié(e) : {nom_salarie_defaut} | Période : {periode_rapport}"
            ws["A2"].font = font_sous_titre
            ws.append([])

            colonnes_excel = [c for c in colonnes_attendues if c != "ID"]
            ws.append(colonnes_excel)
            for col_num in range(1, len(colonnes_excel) + 1):
                cell = ws.cell(row=4, column=col_num)
                cell.font = font_en_tete
                cell.fill = fill_en_tete
                cell.alignment = align_centre
                cell.border = bordure_fine

            for _, row in st.session_state["frais_data"].iterrows():
                valeurs = []
                for c_name in colonnes_excel:
                    val = row[c_name]
                    if c_name in COLS_MONTANTS:
                        try:
                            val = float(val) if pd.notna(val) and val != "" else 0.0
                        except ValueError:
                            val = 0.0
                    elif pd.isna(val):
                        val = ""
                    valeurs.append(val)
                ws.append(valeurs)

                r = ws.max_row
                for col_num, c_name in enumerate(colonnes_excel, start=1):
                    cell = ws.cell(row=r, column=col_num)
                    cell.font = font_donnees
                    cell.border = bordure_fine
                    if c_name in COLS_MONTANTS:
                        cell.number_format = "#,##0.00 €"
                        cell.alignment = align_droite
                    else:
                        cell.alignment = align_gauche

            ligne_totaux = []
            for c_name in colonnes_excel:
                if c_name in COLS_MONTANTS:
                    ligne_totaux.append(
                        pd.to_numeric(
                            st.session_state["frais_data"][c_name], errors="coerce"
                        ).sum()
                    )
                elif c_name == "Libellé":
                    ligne_totaux.append("TOTAL GÉNÉRAL")
                else:
                    ligne_totaux.append("")
            ws.append(ligne_totaux)

            tot_row = ws.max_row
            for col_num, c_name in enumerate(colonnes_excel, start=1):
                cell = ws.cell(row=tot_row, column=col_num)
                cell.font = font_total
                cell.border = bordure_total
                if c_name in COLS_MONTANTS:
                    cell.number_format = "#,##0.00 €"
                    cell.alignment = align_droite
                else:
                    cell.alignment = align_gauche

            for col in ws.columns:
                max_len = max(
                    (len(str(c.value)) for c in col if c.value is not None), default=0
                )
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(
                    max_len + 4, 12
                )

            output_excel = io.BytesIO()
            wb.save(output_excel)
            st.session_state["excel_bytes"] = output_excel.getvalue()
            st.session_state["excel_nom"] = (
                f"NDF-{nom_salarie_defaut.replace(' ', '-')}-{periode_rapport}.xlsx"
            )

            memory_zip = io.BytesIO()
            with zipfile.ZipFile(memory_zip, "w", zipfile.ZIP_DEFLATED) as zf:
                fichiers_actifs = st.session_state["frais_data"][
                    "Justificatif"
                ].tolist()
                for nom_fic in fichiers_actifs:
                    if pd.notna(nom_fic) and nom_fic != "Aucun justificatif":
                        chemin_complet = os.path.join(DOSSIER_JUSTIFICATIFS, nom_fic)
                        if os.path.isfile(chemin_complet):
                            zf.write(chemin_complet, arcname=nom_fic)
            memory_zip.seek(0)
            st.session_state["zip_bytes"] = memory_zip.getvalue()
            st.session_state["zip_nom"] = (
                f"Justificatifs-{nom_salarie_defaut.replace(' ', '-')}-{periode_rapport}.zip"
            )

        if st.session_state.get("excel_bytes"):
            st.success("🎉 Fichiers générés !")
            col_dl1, col_dl2 = st.columns(2)
            with col_dl1:
                st.download_button(
                    label="📥 Télécharger l'Excel",
                    data=st.session_state["excel_bytes"],
                    file_name=st.session_state.get(
                        "excel_nom", "note_de_frais.xlsx"
                    ),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            with col_dl2:
                if st.session_state.get("zip_bytes"):
                    st.download_button(
                        label="📥 Télécharger le ZIP",
                        data=st.session_state["zip_bytes"],
                        file_name=st.session_state.get(
                            "zip_nom", "justificatifs.zip"
                        ),
                        mime="application/zip",
                    )

        if st.button("🧹 Archiver et vider pour un nouveau mois"):
            if not st.session_state["frais_data"].empty:
                periode_actuelle = (
                    str(st.session_state["frais_data"]["Période"].iloc[0])
                    if "Période" in st.session_state["frais_data"].columns
                    else "archive"
                )
                horodatage = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                nom_archive = f"frais_{user_key}_{periode_actuelle}_{horodatage}.csv"
                st.session_state["frais_data"].to_csv(
                    os.path.join(DOSSIER_ARCHIVES, nom_archive), index=False
                )

                if os.path.exists(DOSSIER_JUSTIFICATIFS):
                    for f in os.listdir(DOSSIER_JUSTIFICATIFS):
                        chemin_f = os.path.join(DOSSIER_JUSTIFICATIFS, f)
                        if os.path.isfile(chemin_f):
                            os.remove(chemin_f)

                st.session_state["frais_data"] = pd.DataFrame(columns=colonnes_attendues)
                st.session_state["id_counter"] = 0
                st.session_state.pop("excel_bytes", None)
                st.session_state.pop("zip_bytes", None)
                if os.path.exists(FICHIER_SAUVEGARDE):
                    os.remove(FICHIER_SAUVEGARDE)
                st.rerun()

else:
    st.info("Aucune dépense enregistrée pour le moment.")
